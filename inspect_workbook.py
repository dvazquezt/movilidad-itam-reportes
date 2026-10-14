"""
Inspección estructural del libro (sin modificarlo): mapa de hojas, fórmulas,
referencias #REF!, errores cacheados, dependencias entre hojas y tablas
dinámicas con su rango fuente.  Útil para correr cada vez que llegue una
versión nueva de la "macro" y detectar cambios de estructura.

Uso:  python inspect_workbook.py "datos_macro_actual al 6-10-2026.xlsx" [--out mapa_hojas.xlsx]
"""
from __future__ import annotations

import argparse
import posixpath
import re
import warnings
import zipfile
from collections import Counter
from datetime import date, timedelta

import openpyxl
import pandas as pd

warnings.filterwarnings("ignore")
ERRORS = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}


def pivot_map(path) -> pd.DataFrame:
    z = zipfile.ZipFile(path)
    wb = z.read("xl/workbook.xml").decode()
    rels = z.read("xl/_rels/workbook.xml.rels").decode()
    rid = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
    sheet_file = {posixpath.basename(rid[r]): n for n, r in re.findall(r'<sheet name="([^"]+)"[^>]*r:id="(rId\d+)"', wb)}
    caches = {}
    for cid, r in re.findall(r'<pivotCache cacheId="(\d+)" r:id="(rId\d+)"', wb):
        x = z.read("xl/" + rid[r]).decode()
        src = re.search(r'<worksheetSource ref="([^"]+)" sheet="([^"]+)"', x)
        ref = re.search(r'refreshedDate="([\d.]+)"', x)
        caches[cid] = {
            "rango_fuente": f"{src.group(2)}!{src.group(1)}" if src else "?",
            "actualizado": (date(1899, 12, 30) + timedelta(days=int(float(ref.group(1))))).isoformat() if ref else "",
            "campos": re.findall(r'<cacheField name="([^"]+)"', x),
        }
    rows = []
    for name in z.namelist():
        if not re.match(r"xl/worksheets/_rels/sheet\d+\.xml\.rels", name):
            continue
        sheet = sheet_file.get(posixpath.basename(name).replace(".rels", ""), "?")
        for pt in re.findall(r'Target="\.\./pivotTables/([^"]+)"', z.read(name).decode()):
            x = z.read("xl/pivotTables/" + pt).decode()
            cid = re.search(r'cacheId="(\d+)"', x).group(1)
            fl = caches[cid]["campos"]

            def fields(tag):
                blk = re.search(rf"<{tag}.*?</{tag}>", x, re.S)
                return [fl[int(i)] if int(i) >= 0 else "Σ" for i in re.findall(r'<field x="(-?\d+)"', blk.group(0))] if blk else []
            rows.append({
                "hoja": sheet, "tabla": re.search(r'name="([^"]+)"', x).group(1),
                "rango_fuente": caches[cid]["rango_fuente"], "cache_actualizado": caches[cid]["actualizado"],
                "filas": ", ".join(fields("rowFields")), "columnas": ", ".join(fields("colFields")),
                "filtros": ", ".join(fl[int(i)] for i in re.findall(r'<pageField fld="(\d+)"', x)),
                "valores": ", ".join(f"{n} ({fl[int(f)]})" for n, f in re.findall(r'<dataField name="([^"]+)" fld="(\d+)"', x)),
            })
    return pd.DataFrame(rows)


def sheet_map(path) -> pd.DataFrame:
    wbf = openpyxl.load_workbook(path, read_only=True, data_only=False)
    wbv = openpyxl.load_workbook(path, read_only=True, data_only=True)
    states = {ws.title: ws.sheet_state for ws in openpyxl.load_workbook(path, read_only=True).worksheets}
    rows = []
    for name in wbf.sheetnames:
        nf = nref = nval = maxr = maxc = 0
        deps = Counter()
        for row in wbf[name].iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                nval += 1
                maxr, maxc = max(maxr, c.row), max(maxc, c.column)
                if isinstance(v, str) and v.startswith("="):
                    nf += 1
                    nref += "#REF!" in v
                    for s in re.findall(r"'([^']+)'!|([A-Za-z0-9_\-]+)!", v):
                        deps[(s[0] or s[1])] += 1
        errs = Counter(c.value for row in wbv[name].iter_rows() for c in row
                       if isinstance(c.value, str) and c.value in ERRORS)
        rows.append({"hoja": name, "estado": states.get(name), "filas": maxr, "cols": maxc,
                     "celdas": nval, "fórmulas": nf, "fórmulas_con_#REF": nref,
                     "errores_cacheados": dict(errs) or "", "depende_de": ", ".join(k for k in deps if k != "REF")})
    return pd.DataFrame(rows)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    sm, pm = sheet_map(a.path), pivot_map(a.path)
    pd.set_option("display.width", 220)
    print(sm.to_string(index=False))
    print(pm.to_string(index=False))
    if a.out:
        with pd.ExcelWriter(a.out) as xw:
            sm.to_excel(xw, sheet_name="hojas", index=False)
            pm.to_excel(xw, sheet_name="tablas_dinamicas", index=False)
        print("->", a.out)
