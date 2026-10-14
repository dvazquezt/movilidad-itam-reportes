"""
Carga de datos.  ÚNICO módulo que sabe que la fuente es un Excel.

Cuando haya acceso a Dataverse se agrega `load_from_dataverse()` que devuelva
EXACTAMENTE el mismo DataFrame (mismas columnas que config.INOUT_COLUMNS +
`_fuente_fila`) y el resto del pipeline no cambia.

El archivo original se abre en modo sólo lectura (openpyxl read_only=True);
nunca se escribe sobre él.
"""
from __future__ import annotations

import hashlib
import warnings
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
import openpyxl

import config as C

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


@dataclass
class SourceData:
    inout: pd.DataFrame                       # valores (cacheados) de In-Out
    inout_formulas: pd.DataFrame | None       # mismo shape, texto de fórmula o valor
    unis: pd.DataFrame
    carreras: pd.DataFrame
    pais_region: pd.DataFrame
    raw_sheets: dict = field(default_factory=dict)   # hojas de reporte (valores) para paridad
    raw_formulas: dict = field(default_factory=dict) # hojas de reporte (fórmulas)
    source_path: str = ""
    source_sha256: str = ""


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _sheet_to_grid(ws) -> pd.DataFrame:
    rows = list(ws.iter_rows(values_only=True))
    return pd.DataFrame(rows)


def _grid_to_table(grid: pd.DataFrame, header_row: int = 0) -> pd.DataFrame:
    header = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(grid.iloc[header_row])]
    df = grid.iloc[header_row + 1:].copy()
    df.columns = header
    df = df.dropna(how="all")
    return df


def load_inout(wb_values, wb_formulas=None) -> tuple[pd.DataFrame, pd.DataFrame | None]:
    grid = _sheet_to_grid(wb_values[C.SHEET_INOUT])
    header = [str(h) if h is not None else None for h in grid.iloc[0]]
    missing = [c for c in C.INOUT_COLUMNS if c not in header]
    extra = [c for c in header if c and c not in C.INOUT_COLUMNS]
    if missing:
        raise ValueError(f"Faltan columnas esperadas en {C.SHEET_INOUT}: {missing}")
    if extra:
        print(f"[aviso] columnas nuevas no esperadas en {C.SHEET_INOUT}: {extra}")
    df = grid.iloc[1:].copy()
    df.columns = header
    df["_fuente_fila"] = range(2, len(df) + 2)          # fila real en Excel (trazabilidad)
    data_cols = [c for c in df.columns if c in C.INOUT_COLUMNS]
    df = df[df[data_cols].notna().any(axis=1)].reset_index(drop=True)

    fdf = None
    if wb_formulas is not None:
        fgrid = _sheet_to_grid(wb_formulas[C.SHEET_INOUT])
        fdf = fgrid.iloc[1:].copy()
        fdf.columns = header
        fdf["_fuente_fila"] = range(2, len(fdf) + 2)
        fdf = fdf[fdf["_fuente_fila"].isin(df["_fuente_fila"])].reset_index(drop=True)
    return df, fdf


def load_source(path: str | Path, with_formulas: bool = True,
                report_sheets: tuple = (C.SHEET_BALANCE, C.SHEET_REPORTE_ANUAL, C.SHEET_PIM)) -> SourceData:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    print(f"[load] {path.name}  (sólo lectura)")
    wbv = openpyxl.load_workbook(path, read_only=True, data_only=True)
    wbf = openpyxl.load_workbook(path, read_only=True, data_only=False) if with_formulas else None

    inout, inout_f = load_inout(wbv, wbf)
    unis = _grid_to_table(_sheet_to_grid(wbv[C.SHEET_UNIS]))
    unis = unis[unis["Universidad"].notna()]
    carreras = _grid_to_table(_sheet_to_grid(wbv[C.SHEET_CARRERAS]))
    carreras = carreras[carreras["Clave"].notna()]
    pais = _grid_to_table(_sheet_to_grid(wbv[C.SHEET_PAIS_REGION]))
    pais = pais[pais["PAÍS"].notna()]

    raw, rawf = {}, {}
    for s in report_sheets:
        if s in wbv.sheetnames:
            raw[s] = _sheet_to_grid(wbv[s])
            if wbf is not None:
                rawf[s] = _sheet_to_grid(wbf[s])
    wbv.close()
    if wbf is not None:
        wbf.close()
    print(f"[load] In-Out: {len(inout):,} registros | unis-info: {len(unis)} | "
          f"clave-carrera: {len(carreras)} | país-región: {len(pais)}")
    return SourceData(inout, inout_f, unis, carreras, pais, raw, rawf,
                      str(path), _sha256(path))


def cell(grid: pd.DataFrame, ref: str):
    """Lee una celda tipo 'C1' de una hoja cargada como grid (0-index)."""
    col = 0
    letters = "".join(ch for ch in ref if ch.isalpha())
    row = int("".join(ch for ch in ref if ch.isdigit()))
    for ch in letters.upper():
        col = col * 26 + (ord(ch) - 64)
    try:
        return grid.iat[row - 1, col - 1]
    except IndexError:
        return None
