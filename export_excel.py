"""
Exportación a Excel.  Siempre escribe archivos NUEVOS con fecha en el nombre;
nunca toca el archivo fuente.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _format_sheet(ws, df: pd.DataFrame, index: bool):
    ws.freeze_panes = "A2"
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
    cols = ([df.index.name or ""] if index else []) + [str(c) for c in df.columns]
    for i, name in enumerate(cols, start=1):
        sample = [len(str(name))] + [len(str(v)) for v in (df.iloc[:200, i - 1 - int(index)] if i - 1 - int(index) >= 0 else df.index[:200])]
        ws.column_dimensions[get_column_letter(i)].width = min(max(sample) + 2, 60)
    ws.auto_filter.ref = ws.dimensions


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame], meta: dict | None = None,
                   index_sheets: set[str] = frozenset()):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        if meta:
            m = pd.DataFrame({"campo": list(meta), "valor": [str(v) for v in meta.values()]})
            m.to_excel(xw, sheet_name="LEEME", index=False)
            _format_sheet(xw.sheets["LEEME"], m, False)
            xw.sheets["LEEME"].column_dimensions["B"].width = 110
        for name, df in sheets.items():
            sname = name[:31]
            idx = name in index_sheets
            df.to_excel(xw, sheet_name=sname, index=idx)
            _format_sheet(xw.sheets[sname], df, idx)
    return path


def stamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M")
